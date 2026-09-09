"""
common.py — shared Word-document construction helpers.

Every document in the submission is built through these, so they all share one
typographic system: the same margins, the same heading scale, the same table
look, and the brand amber used only for emphasis.
"""
import os

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

import config as CFG

ROOT = os.path.dirname(os.path.abspath(__file__))


def rgb(hexstr):
    hexstr = hexstr.lstrip("#")
    return RGBColor(int(hexstr[0:2], 16), int(hexstr[2:4], 16), int(hexstr[4:6], 16))


INK = rgb("1F2933")
MUTED = rgb("5B6873")
AMBER = rgb("A8790A")      # the brand amber, darkened for ink on white paper
RULE = "C8CED5"


# ── Low-level XML helpers ───────────────────────────────────────────────
def _shade(cell, hexfill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hexfill)
    tcPr.append(shd)


def _borders(table, colour=RULE, size=4):
    tbl = table._tbl
    tblPr = tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), str(size))
        el.set(qn("w:space"), "0")
        el.set(qn("w:color"), colour)
        borders.append(el)
    tblPr.append(borders)


def _keep_header_visible(row):
    trPr = row._tr.get_or_add_trPr()
    el = OxmlElement("w:tblHeader")
    el.set(qn("w:val"), "true")
    trPr.append(el)


def _field(paragraph, instr):
    run = paragraph.add_run()
    a = OxmlElement("w:fldChar"); a.set(qn("w:fldCharType"), "begin")
    b = OxmlElement("w:instrText"); b.set(qn("xml:space"), "preserve"); b.text = instr
    c = OxmlElement("w:fldChar"); c.set(qn("w:fldCharType"), "end")
    run._r.append(a); run._r.append(b); run._r.append(c)
    return run


# ── Document scaffolding ────────────────────────────────────────────────
def new_doc(footer_text=None):
    doc = Document()

    for s in doc.sections:
        s.top_margin = Cm(2.2)
        s.bottom_margin = Cm(2.0)
        s.left_margin = Cm(2.3)
        s.right_margin = Cm(2.3)

    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = INK
    normal.paragraph_format.space_after = Pt(7)
    normal.paragraph_format.line_spacing = 1.18
    rpr = normal.element.get_or_add_rPr()
    rfonts = rpr.get_or_add_rFonts()
    rfonts.set(qn("w:eastAsia"), "Calibri")

    for level, size, colour, before in (
        (1, 17, AMBER, 18),
        (2, 13, INK, 14),
        (3, 11.5, INK, 11),
    ):
        st = doc.styles[f"Heading {level}"]
        st.font.name = "Calibri"
        st.font.size = Pt(size)
        st.font.bold = True
        st.font.color.rgb = colour
        st.paragraph_format.space_before = Pt(before)
        st.paragraph_format.space_after = Pt(5)
        st.paragraph_format.keep_with_next = True

    footer(doc, footer_text)
    return doc


def footer(doc, text=None):
    text = text or f"{CFG.COMPANY} · {CFG.PRODUCT} · {CFG.MODULE}"
    p = doc.sections[0].footer.paragraphs[0]
    p.text = ""
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(text + "     |     Page ")
    r.font.size = Pt(8)
    r.font.color.rgb = MUTED
    fr = _field(p, " PAGE ")
    fr.font.size = Pt(8)
    fr.font.color.rgb = MUTED


def cover(doc, task, title, subtitle, extra=None):
    """Full cover page: task badge, title, subtitle, then a metadata block."""
    sp = doc.add_paragraph()
    sp.paragraph_format.space_after = Pt(90)

    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(2)
    r = p.add_run(task.upper())
    r.font.size = Pt(11)
    r.font.bold = True
    r.font.color.rgb = AMBER
    r.font.name = "Consolas"

    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(4)
    r = p.add_run(title)
    r.font.size = Pt(30)
    r.font.bold = True
    r.font.color.rgb = INK

    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(34)
    r = p.add_run(subtitle)
    r.font.size = Pt(13)
    r.font.color.rgb = MUTED

    rows = [
        ("Team", f"{CFG.TEAM_NUMBER} — {CFG.COMPANY}"),
        ("Product", CFG.PRODUCT),
        ("Client", CFG.CLIENT_ORG),
        ("Module", f"{CFG.MODULE} — {CFG.MODULE_LEADER}"),
        ("Academic year", CFG.ACADEMIC_YEAR),
        ("Date", CFG.SUBMISSION_DATE),
    ]
    if extra:
        rows.extend(extra)

    t = doc.add_table(rows=0, cols=2)
    t.alignment = WD_TABLE_ALIGNMENT.LEFT
    for k, v in rows:
        c = t.add_row().cells
        rk = c[0].paragraphs[0].add_run(k)
        rk.font.bold = True
        rk.font.size = Pt(9.5)
        rk.font.color.rgb = MUTED
        rv = c[1].paragraphs[0].add_run(v)
        rv.font.size = Pt(9.5)
        c[0].width = Cm(3.6)
        c[1].width = Cm(11.5)
    for row in t.rows:
        for cell in row.cells:
            cell.paragraphs[0].paragraph_format.space_after = Pt(2)

    page_break(doc)


def page_break(doc):
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)


def h1(doc, text):
    return doc.add_heading(text, level=1)


def h2(doc, text):
    return doc.add_heading(text, level=2)


def h3(doc, text):
    return doc.add_heading(text, level=3)


def para(doc, text, size=10.5, italic=False, colour=None, after=7, align=None):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(after)
    if align == "center":
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _rich(p, text, size=size, italic=italic, colour=colour)
    return p


def _rich(paragraph, text, size=10.5, italic=False, colour=None):
    """Minimal inline markup: **bold**, `mono`."""
    import re

    def emit(txt, bold=False, mono=False):
        r = paragraph.add_run(txt)
        r.font.bold = bold
        r.font.italic = italic
        if mono:
            r.font.name = "Consolas"
            r.font.size = Pt(size - 1)
            r.font.color.rgb = rgb("7A4B00")
            return
        r.font.size = Pt(size)
        if colour is not None:
            r.font.color.rgb = colour

    for chunk in re.split(r"(\*\*(?:[^*]|\*(?!\*))+\*\*|`[^`]+`)", text):
        if not chunk:
            continue
        if chunk.startswith("**") and chunk.endswith("**"):
            # A code span can sit inside a bold span; handle the nesting rather
            # than emitting the backticks as literal text.
            for sub in re.split(r"(`[^`]+`)", chunk[2:-2]):
                if not sub:
                    continue
                if sub.startswith("`") and sub.endswith("`"):
                    emit(sub[1:-1], bold=True, mono=True)
                else:
                    emit(sub, bold=True)
        elif chunk.startswith("`") and chunk.endswith("`"):
            emit(chunk[1:-1], mono=True)
        else:
            emit(chunk)
    return paragraph


def bullets(doc, items, style="List Bullet", size=10.5):
    for it in items:
        p = doc.add_paragraph(style=style)
        p.paragraph_format.space_after = Pt(3)
        _rich(p, it, size=size)


def numbered(doc, items, size=10.5):
    bullets(doc, items, style="List Number", size=size)


def table(doc, headers, rows, widths=None, size=9.5, header_fill="1B2530",
          zebra="F4F6F8", align_center=()):
    t = doc.add_table(rows=1, cols=len(headers))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    _borders(t)

    hdr = t.rows[0]
    _keep_header_visible(hdr)
    for i, htext in enumerate(headers):
        cell = hdr.cells[i]
        cell.text = ""
        p = cell.paragraphs[0]
        p.paragraph_format.space_after = Pt(1)
        if i in align_center:
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(str(htext))
        r.font.bold = True
        r.font.size = Pt(size)
        r.font.color.rgb = rgb("FFFFFF")
        _shade(cell, header_fill)

    for n, row in enumerate(rows):
        cells = t.add_row().cells
        for i, val in enumerate(row):
            cell = cells[i]
            cell.text = ""
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(1)
            if i in align_center:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            _rich(p, str(val), size=size)
            if zebra and n % 2 == 1:
                _shade(cell, zebra)

    if widths:
        for row in t.rows:
            for i, w in enumerate(widths):
                row.cells[i].width = Cm(w)

    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return t


def caption(doc, text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(12)
    r = p.add_run(text)
    r.font.size = Pt(8.5)
    r.font.italic = True
    r.font.color.rgb = MUTED


def figure(doc, path, width_cm=16.0, cap=None):
    if not os.path.exists(path):
        para(doc, f"[missing figure: {os.path.basename(path)}]", italic=True, colour=MUTED)
        return
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(8)
    p.paragraph_format.space_after = Pt(2)
    p.add_run().add_picture(path, width=Cm(width_cm))
    if cap:
        caption(doc, cap)


def callout(doc, title, text, fill="FFF6DC", bar="F2B90C"):
    """A boxed note. Used for 'fill this in' markers and for honesty notes."""
    t = doc.add_table(rows=1, cols=1)
    _borders(t, colour=bar, size=8)
    cell = t.rows[0].cells[0]
    cell.text = ""
    _shade(cell, fill)
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(2)
    r = p.add_run(title)
    r.font.bold = True
    r.font.size = Pt(9.5)
    r.font.color.rgb = rgb("6B4A00")
    p2 = cell.add_paragraph()
    p2.paragraph_format.space_after = Pt(2)
    _rich(p2, text, size=9.5)
    doc.add_paragraph().paragraph_format.space_after = Pt(4)
    return t


def todo(doc, text):
    return callout(
        doc,
        "TO COMPLETE BEFORE SUBMISSION",
        text,
        fill="FFF1F1",
        bar="EF4444",
    )


def signature_block(doc, labels=("Signed", "Name", "Date")):
    rows = [[lbl, "…………………………………………………………………"] for lbl in labels]
    table(doc, ["", ""], rows, widths=[3.0, 12.0], size=10)


def save(doc, out_dir, filename):
    os.makedirs(out_dir, exist_ok=True)
    path = os.path.join(out_dir, filename)
    doc.save(path)
    print(f"  · {os.path.relpath(path, os.path.dirname(ROOT))}")
    return path
