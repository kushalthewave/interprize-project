"""
Geometry audit for the deck.

LibreOffice is not available on this machine, so slides cannot be rendered to
images for a visual pass. This checks the defects that a render would catch
geometrically: shapes off the slide, shapes too close to the edge, overlapping
text, and text that is very likely to overflow its box.

Text-fit is an ESTIMATE (character-width heuristic), not a real layout.
"""
import sys
from pptx import Presentation
from pptx.util import Emu

EMU_IN = 914400.0
deck = Presentation(sys.argv[1] if len(sys.argv) > 1 else 'Beat-The-Hazard-Presentation.pptx')
SW = deck.slide_width / EMU_IN
SH = deck.slide_height / EMU_IN
print(f'slide: {SW:.3f} x {SH:.3f} in, {len(deck.slides.__iter__.__self__._sldIdLst)} slides\n')

MARGIN = 0.5
issues = 0

def rect(sh):
    return (sh.left / EMU_IN, sh.top / EMU_IN,
            (sh.left + sh.width) / EMU_IN, (sh.top + sh.height) / EMU_IN)

def overlap(a, b):
    ix = min(a[2], b[2]) - max(a[0], b[0])
    iy = min(a[3], b[3]) - max(a[1], b[1])
    return ix, iy

for i, slide in enumerate(deck.slides, 1):
    probs = []
    texts = []
    for sh in slide.shapes:
        if sh.left is None or sh.width is None:
            continue
        x0, y0, x1, y1 = rect(sh)
        name = (sh.name or 'shape')[:24]
        # full-bleed images and scrims are intentional
        full_bleed = x0 <= 0.02 and y0 <= 0.02 and x1 >= SW - 0.02 and y1 >= SH - 0.02

        if not full_bleed:
            if x0 < -0.01 or y0 < -0.01 or x1 > SW + 0.01 or y1 > SH + 0.01:
                probs.append(f'OFF-SLIDE  {name}: ({x0:.2f},{y0:.2f})-({x1:.2f},{y1:.2f})')
            elif sh.has_text_frame and sh.text_frame.text.strip():
                if x0 < MARGIN - 0.01 or y0 < MARGIN - 0.15 or x1 > SW - MARGIN + 0.01 or y1 > SH - MARGIN + 0.15:
                    probs.append(f'TIGHT MARGIN  {name}: ({x0:.2f},{y0:.2f})-({x1:.2f},{y1:.2f})')

        if sh.has_text_frame and sh.text_frame.text.strip():
            txt = sh.text_frame.text
            sizes = [r.font.size.pt for p in sh.text_frame.paragraphs for r in p.runs if r.font.size]
            pt = max(sizes) if sizes else 14
            w_in = x1 - x0
            h_in = y1 - y0
            # ~0.50 em average glyph advance for Calibri/Arial mixed case
            cpl = max(1, int(w_in * 72 / (pt * 0.50)))
            lines = 0
            for para in txt.split('\n'):
                lines += max(1, -(-len(para) // cpl))
            need = lines * pt * 1.22 / 72.0
            if need > h_in * 1.12:
                probs.append(f'TEXT MAY OVERFLOW  {name}: needs ~{need:.2f}" has {h_in:.2f}"  [{txt[:38]!r}]')
            texts.append((name, (x0, y0, x1, y1), txt[:26]))

    for a in range(len(texts)):
        for b in range(a + 1, len(texts)):
            n1, r1, t1 = texts[a]
            n2, r2, t2 = texts[b]
            ix, iy = overlap(r1, r2)
            if ix > 0.12 and iy > 0.12:
                probs.append(f'TEXT OVERLAP  {t1!r} x {t2!r}  ({ix:.2f}" x {iy:.2f}")')

    if probs:
        issues += len(probs)
        print(f'--- slide {i} ---')
        for p in probs:
            print('   ', p)

print(f'\n{issues} potential issue(s)')
